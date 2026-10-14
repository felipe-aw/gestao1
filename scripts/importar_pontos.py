"""Lê a lista de pontos de automação (aba "Consolidada") e gera um JSON de pontos.

Uso (dentro de gestao/):  python scripts/importar_pontos.py <planilha.xlsx> <saida.json>
O JSON é aplicado no sistema por scripts/aplicar_pontos.mjs (via API, sem parar o servidor).

Colunas da aba Consolidada (1-based): 1 Sist, 2 Pav, 3 Quadro, 4 Controlador, 5 Modelo, 6 Sistema,
7 Ponto, 8 (tag), 9 Equipamento, 10 Descrição, 11 Observação, 12 DI, 13 AI, 14 AO, 15 DO, 16 CS,
17 ETH, 18 KNX, 19 UND, 20 Destino, 21 Cabo, 22 Sensor/Interface, 23 Fornecimento, 24 Ponto futuro,
25 Fornecedor, 26 Compra, 27 Cabo ok, 28 Equipamento ok, 29 Programação ok, 30 Tela ok, 31 Problema.
"""
import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
TIPOS = [(12, "DI"), (13, "AI"), (14, "AO"), (15, "DO"), (16, "CS"), (17, "ETH"), (18, "KNX")]


def txt(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def andar(pav):
    p = (txt(pav) or "").upper()
    if "SS" in p:
        return -1
    m = re.match(r"(\d+)", p)
    return int(m.group(1)) if m else None


def ok(v):
    return isinstance(v, str) and v.strip().lower() == "ok"


def main(arq, saida):
    ws = openpyxl.load_workbook(arq, data_only=True)["Consolidada"]
    # Células mescladas (ex.: Quadro) → propaga o valor para todas as linhas do bloco.
    mescl = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                mescl[(r, c)] = v

    def cel(r, c):
        return mescl.get((r, c), ws.cell(r, c).value)

    pontos = []
    for r in range(3, ws.max_row + 1):
        sist0 = txt(cel(r, 1))
        equip, desc = txt(cel(r, 9)), txt(cel(r, 10))
        if sist0 in (None, "Sist") or not (equip or desc):
            continue  # cabeçalhos repetidos, linhas de controlador e vazias
        tipos = [t for c, t in TIPOS if cel(r, c) not in (None, 0, "0", "")]
        etapas = {k: True for k, c in (("cabo", 27), ("equipamento", 28), ("programacao", 29), ("tela", 30)) if ok(cel(r, c))}
        prob = txt(cel(r, 31))
        pontos.append({
            "linha": r,
            "andar": andar(cel(r, 2)),
            "pav": txt(cel(r, 2)),
            "sistema": txt(cel(r, 6)),
            "ponto": {k: v for k, v in {
                "quadro": txt(cel(r, 3)),
                "controlador": txt(cel(r, 4)),
                "equipamento": equip,
                "descricao": desc,
                "tipoPonto": "/".join(tipos) or None,
                "destino": txt(cel(r, 20)),
                "cabo": txt(cel(r, 21)),
                "sensor": txt(cel(r, 22)),
                "fornecimento": txt(cel(r, 23)),
                "fornecedor": txt(cel(r, 25)),
                "compra": txt(cel(r, 26)),
                "futuro": True if (txt(cel(r, 24)) or "").lower() == "x" else None,
                "problema": None if prob in (None, "X", "x") else prob,
                "linha": r,
            }.items() if v is not None},
            "obs": txt(cel(r, 11)),
            "marcadoX": prob in ("X", "x"),
            "etapas": etapas,
        })
    with open(saida, "w", encoding="utf-8") as f:
        json.dump(pontos, f, ensure_ascii=False, indent=1)
    print(f"{len(pontos)} pontos -> {saida}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
