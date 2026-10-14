"""Gera data/db.json a partir das planilhas Excel da obra.

Uso (dentro de gestao/):  python scripts/importar.py [--forcar]
Lê ../Automacao_Programacao_Quartos.xlsx e ../CFTV_Pavimentos.xlsx.
Sem --forcar, não sobrescreve um db.json existente.
"""
import datetime as dt
import json
import os
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
PASTA_XLSX = os.path.dirname(RAIZ)
DESTINO = os.path.join(RAIZ, "data", "db.json")
AUTO = os.path.join(PASTA_XLSX, "Automacao_Programacao_Quartos.xlsx")
CFTV = os.path.join(PASTA_XLSX, "CFTV_Pavimentos.xlsx")

_seq = 0


def nid(prefixo):
    global _seq
    _seq += 1
    return f"{prefixo}{_seq:05d}"


def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def coluna(ws, col, ini=9):
    out = []
    for r in range(ini, ws.max_row + 1):
        v = txt(ws.cell(r, col).value)
        if v is None:
            break
        out.append(v)
    return out


def limpo(d):
    return {k: v for k, v in d.items() if v is not None and v != ""}


def main():
    if os.path.exists(DESTINO) and "--forcar" not in sys.argv:
        print(f"{DESTINO} já existe. Use --forcar para sobrescrever.")
        return

    wa = openpyxl.load_workbook(AUTO, data_only=True)
    wc = openpyxl.load_workbook(CFTV, data_only=True)
    ca, cc = wa["Config"], wc["Config"]

    inicio = iso(ca["B4"].value) or dt.date.today().isoformat()
    quartos_por_andar = int(ca["B5"].value or 13)

    listas = {
        "statusPendencia": coluna(ca, 1),
        "prioridade": coluna(ca, 2),
        "encontradaEm": ["Instalação", "Programação", "Configuração", "Retorno", "Vistoria AW", "Testes / comissionamento"],
        "locais": [l for l in coluna(ca, 7) if not l.startswith("Q")],
        "categorias": {
            "auto_ac": coluna(ca, 3),
            "auto_q": coluna(ca, 3),
            "cftv": coluna(cc, 3),
            "ca": [
                "Leitor sem comunicação", "Fechadura / eletroímã", "Botoeira / sensor de porta", "Catraca / cancela",
                "Cabeamento", "Rede / comunicação", "Alimentação elétrica", "Cadastro / permissões",
                "Integração (CFTV, elevador, PMS)", "Civil / terceiros", "Outros",
            ],
            "sdai": [
                "Detector com defeito", "Detector faltando", "Acionador manual", "Sirene / sinalizador",
                "Módulo de laço / endereçamento", "Central / programação", "Cabeamento / laço",
                "Alimentação / baterias", "Integração (HVAC, elevador, CA)", "Civil / terceiros", "Outros",
            ],
        },
        "tiposItem": {
            "auto_ac": ["Controlador", "Módulo de relés", "Dimmer", "Sensor de presença", "Teclado / keypad", "Gateway", "Outro"],
            "auto_q": ["Quarto padrão", "Suíte", "PNE", "Outro"],
            "cftv": coluna(cc, 5),
            "ca": ["Leitor de cartão", "Leitor facial", "Fechadura eletromagnética", "Catraca", "Cancela", "Controladora", "Outro"],
            "sdai": ["Detector de fumaça", "Detector de temperatura", "Acionador manual", "Sirene", "Sinalizador audiovisual", "Módulo", "Central", "Outro"],
        },
    }

    config = {
        "obra": "Sofitel",
        "inicioCronograma": inicio,
        "andarMin": 0,
        "andarMax": 25,
        "listas": listas,
    }

    # ---- macros ----
    macros_def = {
        "auto_ac": ["Infraestrutura", "Cabeamento", "Instalação de dispositivos", "Programação áreas comuns", "Testes e integração", "Aprovação AW"],
        "auto_q": ["Infraestrutura", "Instalação de dispositivos", "Programação quartos", "Aprovação AW"],
        "cftv": ["Infraestrutura", "Instalação", "Configuração", "Aprovação AW"],
        "ca": ["Infraestrutura", "Instalação", "Configuração / cadastro", "Testes e integração", "Aprovação AW"],
        "sdai": ["Infraestrutura", "Instalação", "Programação da central", "Testes / comissionamento", "Aprovação AW / Bombeiros"],
    }
    macros = []
    macro_id = {}
    for esc, nomes in macros_def.items():
        for i, n in enumerate(nomes):
            m = {"id": nid("M"), "escopo": esc, "nome": n, "ordem": i + 1}
            macros.append(m)
            macro_id[(esc, n)] = m["id"]

    # ---- atividades (cronogramas das planilhas) ----
    atividades = []
    ws = wa["Cronograma"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] is None:
            continue
        esc = "auto_ac" if r[1] == "Área comum" else "auto_q"
        macro = "Programação áreas comuns" if esc == "auto_ac" else "Programação quartos"
        atividades.append(limpo({
            "id": nid("A"), "macroId": macro_id[(esc, macro)], "nome": r[2], "andar": int(r[0]),
            "responsavelTxt": txt(r[3]), "inicioPrev": iso(r[4]), "fimPrev": iso(r[5]),
            "inicioReal": iso(r[6]), "fimReal": iso(r[7]), "progresso": 0,
        }))
    ws = wc["Cronograma"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0] is None:
            continue
        macro = "Instalação" if r[1] == "Instalação" else "Configuração"
        atividades.append(limpo({
            "id": nid("A"), "macroId": macro_id[("cftv", macro)], "nome": f"{r[1]} CFTV", "andar": int(r[0]),
            "responsavelTxt": txt(r[2]), "inicioPrev": iso(r[3]), "fimPrev": iso(r[4]),
            "inicioReal": iso(r[5]), "fimReal": iso(r[6]), "progresso": 0,
        }))
    for a in atividades:
        a.setdefault("progresso", 0)
        resp = a.pop("responsavelTxt", None)
        if resp:
            a["obs"] = f"Responsável: {resp}"

    # ---- itens ----
    itens = []
    ws = wa["Quartos"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[1]:
            continue
        itens.append(limpo({
            "id": nid("I"), "escopo": "auto_q", "codigo": r[1], "andar": int(r[0]), "tipo": "Quarto padrão",
            "dataConfig": iso(r[3]), "tecnico": txt(r[4]), "achouPendencia": (r[5] == "Sim") if r[5] else None,
            "dataRetorno": iso(r[8]), "nRetornos": r[9] if isinstance(r[9], (int, float)) else None,
            "dataOkAW": iso(r[10]), "aprovadoPor": txt(r[11]), "obs": txt(r[13]),
        }))
    ws = wc["Câmeras"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        extra = " · ".join(x for x in [txt(r[6]), f"canal {r[7]}" if r[7] else None, txt(r[8])] if x)
        obs = " · ".join(x for x in [
            txt(r[22]),
            f"Imagem OK: {r[17]}" if r[17] else None,
            f"Gravação OK: {r[18]}" if r[18] else None,
        ] if x)
        itens.append(limpo({
            "id": nid("I"), "escopo": "cftv", "codigo": r[0], "andar": int(r[1]), "local": txt(r[2]), "tipo": txt(r[3]),
            "modelo": txt(r[4]), "ip": txt(r[5]), "extra": extra, "dataInstalacao": iso(r[9]), "dataConfig": iso(r[10]),
            "tecnico": txt(r[11]), "achouPendencia": (r[12] == "Sim") if r[12] else None, "dataRetorno": iso(r[15]),
            "nRetornos": r[16] if isinstance(r[16], (int, float)) else None, "dataOkAW": iso(r[19]),
            "aprovadoPor": txt(r[20]), "obs": obs,
        }))
    # Lista oficial de quartos por andar (enviada pelo cliente em 08/10/2026 — 170 quartos).
    # Quartos da planilha que não estão nesta lista são descartados.
    cheio = list(range(1, 13))  # 01–12
    quartos_oficiais = {
        20: [1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 13],
        19: [3, 4, 5, 6, 7, 10, 12, 13],
        18: [1, 2, 3, 4, 5, 6, 7, 10, 12, 13],
        17: [1, 2, 3, 4, 5, 6, 7, 10, 12, 13],
        16: [1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 13],
        15: [1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 13],
        14: [1, 2, 3, 4, 6, 7, 8, 9, 10, 12, 13],
        13: [1, 2, 3, 4, 6, 7, 8, 9, 10, 12],
        12: [1, 2, 3, 4, 6, 7, 8, 9, 10, 12],
        11: [1, 2, 3, 4, 6, 7, 8, 9, 10, 12],
        10: cheio, 9: cheio, 8: cheio, 7: cheio, 6: cheio,
        5: list(range(1, 9)),
    }
    validos = {f"Q{a}{n:02d}" for a, ns in quartos_oficiais.items() for n in ns}
    itens = [i for i in itens if i["escopo"] != "auto_q" or i["codigo"] in validos]
    for a, ns in sorted(quartos_oficiais.items()):
        for n in ns:
            cod = f"Q{a}{n:02d}"
            if not any(i["codigo"] == cod for i in itens):
                itens.append({"id": nid("I"), "escopo": "auto_q", "codigo": cod, "andar": a, "tipo": "Quarto padrão"})

    # ---- empresas (a partir dos responsáveis das pendências) ----
    empresas = []

    def empresa(nome, esc):
        for e in empresas:
            if e["nome"] == nome:
                if esc not in e["escopos"]:
                    e["escopos"].append(esc)
                return e["id"]
        e = {"id": nid("E"), "nome": nome, "escopos": [esc]}
        empresas.append(e)
        return e["id"]

    # ---- pendências ----
    pendencias = []
    num = 0
    ws = wa["Pendências"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[6]:
            continue
        num += 1
        local = txt(r[3])
        is_q = bool(local and local.startswith("Q") and local[1:].isdigit())
        esc = "auto_q" if is_q else "auto_ac"
        resp = txt(r[8])
        pendencias.append(limpo({
            "id": nid("P"), "num": num, "data": iso(r[1]), "escopo": esc, "andar": r[2],
            "local": None if is_q else local, "itemCodigo": local if is_q else None,
            "encontradaEm": txt(r[4]), "categoria": txt(r[5]), "descricao": txt(r[6]), "prioridade": txt(r[7]),
            "empresaId": empresa(resp, esc) if resp else None, "prazo": iso(r[9]), "status": txt(r[10]) or "Aberta",
            "dataRegularizacao": iso(r[11]), "obs": txt(r[13]),
        }))
    ws = wc["Pendências"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[7]:
            continue
        num += 1
        pendencias.append(limpo({
            "id": nid("P"), "num": num, "data": iso(r[1]), "escopo": "cftv", "andar": r[2], "itemCodigo": txt(r[3]),
            "local": txt(r[4]), "encontradaEm": txt(r[5]), "categoria": txt(r[6]), "descricao": txt(r[7]),
            "prioridade": txt(r[8]), "responsavel": txt(r[9]), "prazo": iso(r[10]), "status": txt(r[12]) or "Aberta",
            "dataRegularizacao": iso(r[13]), "obs": txt(r[15]),
        }))

    db = {
        "config": config,
        "empresas": empresas,
        "equipes": [],
        "macros": macros,
        "atividades": atividades,
        "itens": itens,
        "pendencias": pendencias,
    }
    os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
    with open(DESTINO, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=1)
    print(
        f"OK -> {DESTINO}\n  macros={len(macros)} atividades={len(atividades)} itens={len(itens)} "
        f"pendencias={len(pendencias)} empresas={len(empresas)}"
    )


if __name__ == "__main__":
    main()
